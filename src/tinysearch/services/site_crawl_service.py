import asyncio
import re
import sys
import tempfile
from collections.abc import Mapping
from contextlib import asynccontextmanager, suppress
from typing import Any
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from docx import Document
from pypdf import PdfReader
from rank_bm25 import BM25Okapi

from tinysearch.services.text_chunking_service import chunk_text
from tinysearch.services.token_counter_service import (
    decode_tokens,
    encode_tokens,
    token_count,
)
from tinysearch.telemetry import span_scope

BOILERPLATE_EXCLUDED_TAGS: list[str] = ["nav", "header", "footer", "aside"]


def ensure_utf8_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


def _truncate_to_max_tokens(
    text: str,
    max_return_tokens: int | None,
    encoding_name: str,
) -> str:
    if max_return_tokens is None:
        return text
    tokens = encode_tokens(text, encoding_name)
    if len(tokens) <= max_return_tokens:
        return text
    return decode_tokens(tokens[:max_return_tokens], encoding_name)


def _pick_markdown_for_chunking(
    markdown_raw: str,
    markdown_fit: str,
    fit_min_chars: int,
) -> tuple[str, str]:
    """Prefer fitted text when supplied, otherwise use the raw representation."""
    fit_stripped = markdown_fit.strip()
    if len(fit_stripped) >= fit_min_chars:
        return fit_stripped, "fit"
    return markdown_raw.strip(), "raw"


async def _accessibility_text(page: Any) -> str:
    """Return Playwright AI accessibility text, with visible-text fallback."""
    body = page.locator("body")
    try:
        snapshot = await body.aria_snapshot(mode="ai")
    except (AttributeError, TypeError):
        snapshot = ""
    if snapshot and snapshot.strip():
        return snapshot.strip()
    return (await body.inner_text()).strip()


class DirectPlaywrightCrawler:
    """Reusable renderer exposing only the crawler contract TinySearch needs."""

    def __init__(self, config: Mapping[str, Any] | None = None) -> None:
        self._config = dict(config or {})
        self._playwright: Any = None
        self._browser: Any = None
        self._owns_browser = True

    async def start(self) -> None:
        if self._browser is not None:
            return
        from playwright.async_api import async_playwright

        self._playwright = await async_playwright().start()
        cdp_url = str(self._config.get("browser_cdp_url") or "").strip()
        if cdp_url:
            self._browser = await self._playwright.chromium.connect_over_cdp(cdp_url)
            self._owns_browser = False
        else:
            self._browser = await self._playwright.chromium.launch(headless=True)
            self._owns_browser = True

    async def close(self) -> None:
        browser, self._browser = self._browser, None
        playwright, self._playwright = self._playwright, None
        if browser is not None and self._owns_browser:
            with suppress(Exception):
                await browser.close()
        if playwright is not None:
            with suppress(Exception):
                await playwright.stop()

    async def __aenter__(self) -> "DirectPlaywrightCrawler":
        await self.start()
        return self

    async def __aexit__(self, exc_type: Any, exc: Any, tb: Any) -> None:
        await self.close()

    def _storage_state_path(self) -> str | None:
        raw = str(self._config.get("browser_storage_state_path") or "").strip()
        if not raw:
            return None
        from pathlib import Path

        path = Path(raw)
        return str(path) if path.is_file() else None

    async def arun(self, *, url: str, config: Any = None) -> dict[str, Any]:
        del config
        await self.start()
        if self._browser is None:
            raise RuntimeError("Playwright browser failed to start")

        context_options: dict[str, Any] = {"locale": "en-US"}
        storage_state = self._storage_state_path()
        if storage_state is not None:
            context_options["storage_state"] = storage_state

        context = await self._browser.new_context(**context_options)
        try:
            page = await context.new_page()
            response = await page.goto(url, wait_until="domcontentloaded")
            html = await page.content()
            text = await _accessibility_text(page)
            try:
                title = await page.title()
            except Exception:
                title = ""
            return {
                "url": page.url or url,
                "redirected_url": page.url or url,
                "html": html,
                "markdown_raw": text,
                "markdown_fit": "",
                "metadata": {
                    "title": title,
                    "status": getattr(response, "status", None) if response else None,
                },
            }
        finally:
            await context.close()


def create_browser_crawler(config: Mapping[str, Any] | None = None) -> DirectPlaywrightCrawler:
    """Construct the reusable direct-Playwright renderer used by scrape."""
    return DirectPlaywrightCrawler(config)


class BrowserCrawlerSession:
    """Lazily reuse one direct Playwright browser across server requests."""

    def __init__(self) -> None:
        self._lock = asyncio.Lock()
        self._crawler: DirectPlaywrightCrawler | None = None
        self._active_leases = 0
        self._idle_task: asyncio.Task[None] | None = None
        self._idle_seconds = 300.0

    @property
    def started(self) -> bool:
        return self._crawler is not None

    def _cancel_idle_task(self) -> None:
        task, self._idle_task = self._idle_task, None
        if task is not None and task is not asyncio.current_task():
            task.cancel()

    async def _close_crawler(self, crawler: DirectPlaywrightCrawler | None) -> None:
        if crawler is None:
            return
        with suppress(Exception):
            await crawler.close()

    async def _close_after_idle(self) -> None:
        try:
            await asyncio.sleep(self._idle_seconds)
            async with self._lock:
                if self._active_leases:
                    return
                crawler, self._crawler = self._crawler, None
                self._idle_task = None
            await self._close_crawler(crawler)
        except asyncio.CancelledError:
            return

    @asynccontextmanager
    async def lease(self, config: Mapping[str, Any]):
        async with self._lock:
            self._cancel_idle_task()
            self._idle_seconds = float(
                config.get("browser_idle_shutdown_seconds") or 300.0
            )
            if self._crawler is None:
                crawler = create_browser_crawler(config)
                try:
                    await crawler.start()
                except Exception:
                    await self._close_crawler(crawler)
                    raise
                self._crawler = crawler
            self._active_leases += 1
            crawler = self._crawler
        try:
            yield crawler
        finally:
            async with self._lock:
                self._active_leases = max(0, self._active_leases - 1)
                if self._active_leases == 0 and self._crawler is not None:
                    self._idle_task = asyncio.create_task(self._close_after_idle())

    async def close(self) -> None:
        task = self._idle_task
        self._cancel_idle_task()
        if task is not None and task is not asyncio.current_task():
            with suppress(asyncio.CancelledError):
                await task
        async with self._lock:
            crawler, self._crawler = self._crawler, None
            self._active_leases = 0
        await self._close_crawler(crawler)


def url_path_suffix(url: str) -> str:
    return urlparse(url).path.lower().rsplit(".", 1)[-1] if "." in urlparse(url).path else ""


def is_document_url(url: str) -> bool:
    return url_path_suffix(url) in {"pdf", "docx", "doc"}


def _download_url_bytes(url: str, *, timeout_seconds: float = 30.0) -> bytes:
    req = Request(
        url,
        headers={
            "User-Agent": "TinySearch/0.1",
            "Accept": "*/*",
        },
    )
    with urlopen(req, timeout=timeout_seconds) as resp:
        return resp.read()


def _extract_pdf_text(data: bytes) -> str:
    with tempfile.SpooledTemporaryFile(max_size=20 * 1024 * 1024) as tmp:
        tmp.write(data)
        tmp.seek(0)
        reader = PdfReader(tmp)
        pages: list[str] = []
        for idx, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"## Page {idx}\n\n{text}")
        return "\n\n".join(pages).strip()


def _extract_docx_text(data: bytes) -> str:
    with tempfile.SpooledTemporaryFile(max_size=20 * 1024 * 1024) as tmp:
        tmp.write(data)
        tmp.seek(0)
        document = Document(tmp)
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts).strip()


def extract_document_text(url: str, *, timeout_seconds: float = 30.0) -> tuple[str, str]:
    suffix = url_path_suffix(url)
    if suffix == "doc":
        raise ValueError("legacy .doc files are not supported; use PDF or DOCX")

    data = _download_url_bytes(url, timeout_seconds=timeout_seconds)
    if suffix == "pdf":
        return _extract_pdf_text(data), "pdf"
    if suffix == "docx":
        return _extract_docx_text(data), "docx"

    raise ValueError(f"unsupported document type: {suffix or 'unknown'}")


def _tokenize_for_bm25(text: str) -> list[str]:
    return re.findall(r"[a-zA-Z0-9_./:#-]+", text.lower())


def rank_chunks_bm25(
    query: str,
    chunks: list[dict],
    top_k: int = 5,
) -> list[dict]:
    if not query or not chunks:
        return []

    corpus = [_tokenize_for_bm25(chunk["text"]) for chunk in chunks]
    bm25 = BM25Okapi(corpus)
    query_tokens = _tokenize_for_bm25(query)
    scores = bm25.get_scores(query_tokens)
    ranked = sorted(zip(chunks, scores), key=lambda item: item[1], reverse=True)

    return [
        {
            "chunk_id": chunk["chunk_id"],
            "heading": chunk["heading"],
            "score": float(score),
            "tokens": chunk["tokens"],
            "text": chunk["text"],
        }
        for chunk, score in ranked[:top_k]
    ]


def _result_field(result: Any, key: str, default: Any = "") -> Any:
    if isinstance(result, Mapping):
        return result.get(key, default)
    return getattr(result, key, default)


async def fetch_html_for_query(
    url: str,
    user_query: str | None,
    *,
    crawler: Any | None = None,
) -> dict[str, Any]:
    """Render with Playwright and return accessibility text plus full HTML."""
    del user_query
    ensure_utf8_stdio()

    with span_scope(
        "tinysearch.browser",
        attributes={"tinysearch.browser.used": True},
        operation="browser",
    ) as browser_telemetry:
        if crawler is not None:
            result = await crawler.arun(url=url, config=None)
        else:
            async with create_browser_crawler() as owned_crawler:
                result = await owned_crawler.arun(url=url, config=None)
        browser_telemetry.complete()

    metadata_obj = _result_field(result, "metadata", {})
    metadata = dict(metadata_obj) if isinstance(metadata_obj, Mapping) else {}
    final_url = (
        _result_field(result, "redirected_url", None)
        or _result_field(result, "url", None)
        or url
    )

    return {
        "final_url": str(final_url),
        "html": str(_result_field(result, "html", "") or ""),
        "markdown_raw": str(_result_field(result, "markdown_raw", "") or ""),
        "markdown_fit": "",
        "metadata": metadata,
    }


async def crawl(
    url: str,
    max_return_tokens: int = 100000,
    encoding_name: str = "o200k_base",
    *,
    user_query: str | None = None,
    crawler: Any | None = None,
) -> dict:
    page = await fetch_html_for_query(
        url,
        user_query,
        crawler=crawler,
    )
    markdown_raw = _truncate_to_max_tokens(
        page["markdown_raw"], max_return_tokens, encoding_name
    )
    markdown_body, markdown_source = markdown_raw.strip(), "raw"
    return {
        "url": page["final_url"],
        "max_return_tokens": max_return_tokens,
        "html": page["html"],
        "markdown_raw": markdown_raw,
        "markdown": markdown_body,
        "markdown_fit": "",
        "markdown_source": markdown_source,
        "tokens_raw": token_count(markdown_raw, encoding_name),
    }


async def crawl_search(
    url: str,
    user_query: str,
    top_k: int = 5,
    max_chunk_tokens: int = 500,
    overlap_tokens: int = 80,
    max_return_tokens: int | None = None,
    encoding_name: str = "o200k_base",
) -> dict:
    """Legacy helper: fetch once, then run TinySearch local BM25 over chunks."""
    ensure_utf8_stdio()

    if is_document_url(url):
        markdown_raw, document_type = await asyncio.to_thread(extract_document_text, url)
        html = ""
    else:
        page = await fetch_html_for_query(
            url=url,
            user_query=user_query,
        )
        markdown_raw = page["markdown_raw"]
        html = page["html"]
        document_type = None

    chunks = chunk_text(
        text=markdown_raw,
        max_chunk_tokens=max_chunk_tokens,
        overlap_tokens=overlap_tokens,
        encoding_name=encoding_name,
    )
    ranked_chunks = rank_chunks_bm25(query=user_query, chunks=chunks, top_k=top_k)

    if max_return_tokens is not None:
        for chunk in ranked_chunks:
            tokens = encode_tokens(chunk["text"], encoding_name)
            if len(tokens) > max_return_tokens:
                chunk["text"] = decode_tokens(tokens[:max_return_tokens], encoding_name)
                chunk["tokens"] = max_return_tokens

    result = {
        "url": url,
        "query": user_query,
        "html": html,
        "markdown_raw": markdown_raw,
        "markdown_fit": "",
        "tokens_raw": token_count(markdown_raw, encoding_name),
        "tokens_fit": 0,
        "chunks_total": len(chunks),
        "chunks": chunks,
        "ranked_chunks": ranked_chunks,
    }
    if document_type is not None:
        result["document_type"] = document_type
    return result
